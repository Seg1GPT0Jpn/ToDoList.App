# 全問題の一覧（解答・解説つき）を Excel にする。
#   cd rpg_game
#   dart run tool/dump_questions.dart /tmp/all.json
#   python3 tool/question_list_xlsx.py /tmp/all.json ../docs/question_list.xlsx
import json, sys, datetime, collections, re
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, Alignment, PatternFill

src, dst = sys.argv[1], sys.argv[2]
qs = json.load(open(src))
NAMES = [('math', '数学'), ('science', '理科'), ('social', '社会'), ('japanese', '国語'),
         ('english', '英語'), ('sea', '英文法'), ('information', '情報'), ('music', '音楽')]
ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

wb = Workbook(write_only=True)
bold = Font(bold=True)
head_fill = PatternFill('solid', fgColor='DDEBF7')
wrap = Alignment(wrap_text=True, vertical='top')


def cell(ws, v, header=False, wrapped=False):
    if isinstance(v, str):
        v = ILLEGAL.sub('', v)
    c = WriteOnlyCell(ws, value=v)
    if isinstance(v, str) and v.startswith('='):
        c.data_type = 's'  # 数式にしない
    if header:
        c.font = bold
        c.fill = head_fill
    if wrapped:
        c.alignment = wrap
    return c


by = collections.defaultdict(list)
for q in qs:
    by[q['world']].append(q)

ws = wb.create_sheet('概要')
ws.column_dimensions['A'].width = 16
for col in 'BCD':
    ws.column_dimensions[col].width = 14
ws.append([cell(ws, f'つづりクエスト 全問題一覧（{datetime.date.today():%Y-%m-%d}）', header=True)])
ws.append([])
ws.append([cell(ws, x, header=True) for x in ['教科', '問題数', 'うち収録問題', 'うち自動で作る問題']])
for w, name in NAMES:
    xs = by.get(w, [])
    g = sum(1 for q in xs if q['generated'])
    ws.append([name, len(xs), len(xs) - g, g])
total_g = sum(1 for q in qs if q['generated'])
ws.append([cell(ws, '合計', header=True), len(qs), len(qs) - total_g, total_g])
ws.append([])
for line in [
    '・「収録」は問題ファイルに入っている問題。「自動」はアプリが読みこむときに計算や表から作る問題（毎回同じ問題になる）。',
    '・自動の問題の答えはプログラムが計算する。解説には計算の手順や理由を入れている。',
    '・解説の列には、解説・考え方・よくあるまちがい・段階問題の各問の解説をまとめて入れている。',
    '・正解番号は4択の問題だけ（選択肢1〜4のどれが正解か）。',
    '・市販教材（LEAP など）や教科書の本文は入れていない。',
]:
    ws.append([line])

COLS = [('ID', 22), ('セット', 26), ('単元', 40), ('形式', 8), ('分類', 8), ('問題文', 50), ('本文', 40),
        ('正解', 28), ('選択肢1', 22), ('選択肢2', 22), ('選択肢3', 22), ('選択肢4', 22), ('正解番号', 8),
        ('解説', 60), ('作り方', 8)]
for w, name in NAMES:
    xs = by.get(w, [])
    if not xs:
        continue
    ws = wb.create_sheet(name)
    ws.freeze_panes = 'A2'
    for i, (_, width) in enumerate(COLS):
        ws.column_dimensions[chr(ord('A') + i)].width = width
    ws.auto_filter.ref = f'A1:O{len(xs) + 1}'
    ws.append([cell(ws, h, header=True) for h, _ in COLS])
    for q in xs:
        ch = (q['choices'] + ['', '', '', ''])[:4] if len(q['choices']) <= 4 else q['choices'][:3] + [' / '.join(q['choices'][3:])]
        row = [q['id'], q['set'], q['unit'], q['format'], q['category'], q['prompt'], q['body'], q['answer'],
               *ch, q['answerNo'], q['explanation'], '自動' if q['generated'] else '収録']
        ws.append([cell(ws, v, wrapped=(i in (5, 6, 13))) for i, v in enumerate(row)])
wb.save(dst)
print('saved', dst, len(qs))
